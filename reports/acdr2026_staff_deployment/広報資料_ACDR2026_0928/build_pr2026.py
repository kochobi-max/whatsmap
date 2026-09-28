from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
import copy
def setp(p,text):
    for h in list(p._p):
        if h.tag.endswith('}hyperlink'): p._p.remove(h)
    rs=p.runs
    if not rs: p.add_run(text); return
    rs[0].text=text
    for r in rs[1:]: r._r.getparent().remove(r._r)
def strip_bm(e):
    for x in list(e.iter()):
        if x.tag.endswith('}bookmarkStart') or x.tag.endswith('}bookmarkEnd'): x.getparent().remove(x)
    return e
def ins_after(p,text):
    e=strip_bm(copy.deepcopy(p._p)); p._p.addnext(e); q=Paragraph(e,p._parent); setp(q,text); return q
# ================= A. 内閣府 記者発表 =================
d=Document('b_2025_acc.docx'); P=d.paragraphs
setp(P[0],'令和８年10月　日')
setp(P[6],'「アジア防災会議2026」をシンガポールで開催\n～アジア各国の防災能力の向上とアジア地域の防災ネットワークの充実・強化を図ります～')
setp(P[10],'今回、シンガポール民間防衛庁（SCDF）の協力の下、シンガポール国際防災・危機管理エキスポ（SIDEX）と合同で「SIDEX-ACDR2026」として、シンガポールEXPOにてアジア防災会議2026（ACDR2026）が開催され、10月28日（水）の開会式には、内閣府から〇〇が出席（予定）します。')
setp(P[11],'なお、ACDR2026ウェビナー（URLは確定後に記載）より、どなたでも会議を御覧いただけます。')
setp(P[13],'＜アジア防災会議2026の概要＞')
setp(P[15],'１　日程　　令和８年10月28日（水）、29日（木）　※プログラムは別紙ご参照')
ins_after(P[15],'　　　　　　　　（SIDEX-ACDR2026全体は10月27日（火）から30日（金）まで）')
setp(P[16],'２　場所  　シンガポールEXPO（シンガポール）')
setp(P[18],'　　　　　　　　（ACDR2026ウェビナー（URLは確定後に記載））')
setp(P[19],'３　主催　　内閣府政策統括官（防災担当）、アジア防災センター（ADRC）')
ins_after(P[19],'　　　　　　　　（SIDEX主催：COSEM（シンガポール民間防衛庁職員協同組合）、協力：シンガポール民間防衛庁（SCDF））')
setp(P[22],'５　使用言語　　英語　')
setp(P[23],'６　特設HP　　https://acdr.adrc.asia/home/acdr2026　※会議の詳細内容はこちら')
P[30]._p.getparent().remove(P[30]._p)
P[14]._p.getparent().remove(P[14]._p)
P[8]._p.getparent().remove(P[8]._p)
P[12]._p.getparent().remove(P[12]._p)
d.save('out_内閣府記者発表資料_ACDR2026_案_0928.docx')
# ================= B. ADRC 記者発表（資料配付） =================
d=Document('d_acdr2025r_acc.docx'); P=d.paragraphs
t=d.tables[0]; r=t.rows[2]
def setcell(c,text):
    ps=c.paragraphs; setp(ps[0],text)
    for p in ps[1:]: p._p.getparent().remove(p._p)
setcell(r.cells[0],'10/　(　)\n　：00'); setcell(r.cells[3],'理事長　小川　雄二郎')
setp(P[3],'「アジア防災会議2026」の開催について')
setp(P[7],'今回、シンガポール民間防衛庁（SCDF）の協力の下、シンガポール国際防災・危機管理エキスポ（SIDEX）と合同で「SIDEX-ACDR2026」として、シンガポールEXPOにてアジア防災会議2026（ACDR2026）が開催され、10月28日（水）の開会式には、内閣府から〇〇が出席（予定）します。')
setp(P[8],'なお、ACDR2026ウェビナー（URLは確定後に記載）より、どなたでも会議を御覧いただけます。')
setp(P[12],'１　開 催 日　　2026年10月28日（水）、29日（木）')
ins_after(P[12],'　　　　　　　　（SIDEX-ACDR2026全体は10月27日（火）から30日（金）まで）')
setp(P[13],'２　場　　所  　シンガポールEXPO（シンガポール）')
setp(P[15],'（ACDR2026ウェビナー（URLは確定後に記載））')
setp(P[17],'３　主　　催　　内閣府政策統括官（防災担当）、一般財団法人アジア防災センター')
ins_after(P[17],'　　　　　　　　（SIDEX主催：COSEM（シンガポール民間防衛庁職員協同組合）、協力：シンガポール民間防衛庁（SCDF））')
setp(P[21],'　〇10月28日（水）　')
setp(P[22],'・合同開会式（SIDEX-ACDR2026）　シンガポール政府代表、日本政府代表　他')
setp(P[23],'・ACDR2026開会挨拶　日本政府代表、シンガポール民間防衛庁長官、アジア防災センター長')
setp(P[24],'・ラウンドテーブル（第１部）　仙台防災枠組の実施状況と残る課題に関するメンバー国からの発表')
setp(P[25],'・アジア防災センター運営委員会（メンバー国のみ）')
setp(P[26],'　〇10月29日（木）　')
setp(P[27],'・ラウンドテーブル（第２部）　仙台防災枠組の実施状況と残る課題に関するメンバー国からの発表')
q=ins_after(P[27],'・テクニカルセッション　孤立地域・アクセス困難地域の防災DX：先端技術によるコミュニティ・レジリエンスの再定義')
q=ins_after(q,'・フィールドトリップ　シンガポール民間防衛アカデミー（CDA）')
q=ins_after(q,'　〇10月30日（金）　')
q=ins_after(q,'・SIDEX-ACDR2026会議（基調講演、パネルディスカッション）')
setp(P[29],'６　使用言語　　英語　')
setp(P[30],'７　特設ＨＰ　　https://acdr.adrc.asia/home/acdr2026'); P[30].alignment=0
for t in d.element.body.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'):
    if t.text and 'acdr2025' in t.text: t.text=t.text.replace('acdr2025','acdr2026')
d.save('out_記者発表（資料配付）_ACDR2026_案_0928.docx')
# ================= C. 別紙１ プログラム =================
d=Document('c_jpn_agenda_acc.docx'); P=d.paragraphs
setp(P[0],'更新：0928')
setp(P[2],'アジア防災会議　２０２６')
setp(P[3],'先端技術で未来に備える越境的な気候・災害リスク削減 ― 強靭な社会に向けて（仮訳）')
q=ins_after(P[3],'（SIDEX-ACDR2026全体テーマ：災害に強い未来 ― リーダーシップ、連帯、備え）')
setp(P[4],'2026年10月27日－30日')
setp(P[5],'会場：シンガポールEXPO（シンガポール）')
body=d.element.body
tbls=[t for t in body if t.tag.endswith('}tbl')]
tpl_tbl=copy.deepcopy(tbls[0])
tpl_row=copy.deepcopy(Table(tbls[0],d).rows[0]._tr)
tpl_head=copy.deepcopy(P[9]._p)   # '１日目: ...'
t0=Table(tbls[0],d); c1=t0.rows[0].cells[1]
tpl_title=copy.deepcopy(c1.paragraphs[1]._p)  # bold single run
tpl_body=copy.deepcopy(c1.paragraphs[2]._p)   # normal single run
tpl_time=copy.deepcopy(t0.rows[0].cells[0].paragraphs[0]._p)
# remove old day headings + tables (everything after P[8] up to sectPr)
els=list(body)
start=els.index(P[8]._p)+1
for el in els[start:]:
    if el.tag.endswith('}sectPr'): continue
    body.remove(el)
sect=[el for el in body if el.tag.endswith('}sectPr')][0]
days=[
 ('０日目：10月27日（火）',[('19:00-21:00',['SJ60記念レセプション（SIDEX-ACDR2026）','会場：在シンガポール日本国大使公邸','＊招待者のみ'])]),
 ('１日目：10月28日（水）',[
   ('10:00-10:30',['合同開会式（SIDEX-ACDR2026）　於：シンガポールEXPO Hall 3','主賓：シンガポール政府代表（予定）','国家消防・緊急事態準備評議会（NFEC）会長（予定）','日本政府代表（内閣府）']),
   ('10:30-11:15',['VIPツアー（日本パビリオンを含むSIDEX展示の視察）','ACDR参加者：日本パビリオンおよびSIDEX展示の見学（アジア防災センター案内）']),
   ('11:45-12:45',['昼食　於：Hall 3']),
   ('12:45-13:00',['ACDR2026開会式　於：Peridot Room','開会挨拶','シンガポール民間防衛庁（SCDF）長官（予定）','日本政府代表（内閣府）','三浦　房紀　アジア防災センター　センター長']),
   ('13:00-15:00',['ラウンドテーブル・セッション（第１部）：仙台防災枠組2015-2030の実施状況と残る課題','モデレーター：シンガポール民間防衛庁（SCDF）','共同モデレーター：ADRCメンバー国','スピーカー：ADRCメンバー国政府代表']),
   ('15:00-15:15',['休憩']),
   ('15:15-16:15',['アジア防災センター運営委員会（メンバー国のみ）']),
   ('18:30-21:00',['SIDEXガラディナー　於：オーチャードホテル','＊ADRCメンバー国・アドバイザー国・リソース・その他招待者のみ'])]),
 ('２日目：10月29日（木）',[
   ('09:00-09:30',['ラウンドテーブル・セッション（第２部）：仙台防災枠組2015-2030の実施状況と残る課題','モデレーター：シンガポール民間防衛庁（SCDF）','共同モデレーター：ADRCメンバー国']),
   ('09:30-11:30',['テクニカルセッション：孤立地域・アクセス困難地域の防災DX ― 先端技術によるコミュニティ・レジリエンスの再定義','モデレーター：シンガポール民間防衛庁（SCDF）','共同議長：アジア防災センター','スピーカー：','Google（シンガポール）','インドネシア国家研究イノベーション庁（BRIN）','ユース・イノベーション・ラボ（ネパール）','宇宙航空研究開発機構（JAXA）','日本の民間企業（内閣府推薦、調整中）','ウズベキスタン（調整中）']),
   ('11:30-11:45',['休憩']),
   ('11:45-12:00',['まとめ・閉会挨拶','アジア防災センター']),
   ('12:00-13:00',['昼食　於：Hall 3']),
   ('13:00-16:15',['フィールドトリップ：シンガポール民間防衛アカデミー（CDA）','＊メンバー国・リソース登録者・その他招待者のみ、要事前予約']),
   ('18:00-20:05\n20:05-21:30',['SCDFパレード　於：Hall 4','パレード・ディナーレセプション　於：Garnet Room','＊メンバー国・リソース登録者・その他招待者のみ'])]),
 ('３日目：10月30日（金）',[
   ('09:00-12:00',['SIDEX-ACDR2026　基調講演・パネル１　於：Hall 3','基調講演：集団的即応態勢の構築 ― 災害・危機管理における連帯とリーダーシップ','パネル１：災害管理における結束 ― 責任の共有を育む']),
   ('12:00-13:45',['昼食　於：Hall 3']),
   ('13:50-14:50',['展示ツアー']),
   ('15:00-16:30',['パネル２：未来に備える危機管理 ― 危機計画・リスク削減・対応・復旧の強化']),
   ('16:30-16:50',['閉会式'])]),
]
def mkpara(tpl,text):
    e=strip_bm(copy.deepcopy(tpl)); p=Paragraph(e,None); setp(p,text); return e
for head,rows in days:
    h=strip_bm(copy.deepcopy(tpl_head)); sect.addprevious(h); setp(Paragraph(h,None),head)
    tb=strip_bm(copy.deepcopy(tpl_tbl)); sect.addprevious(tb); T=Table(tb,d)
    for tr in list(T._tbl.tr_lst): T._tbl.remove(tr)
    for tm,lines in rows:
        tr=strip_bm(copy.deepcopy(tpl_row)); T._tbl.append(tr)
        row=T.rows[-1]
        c0=row.cells[0]; c1=row.cells[1]
        for c in (c0,c1):
            for p in c.paragraphs: p._p.getparent().remove(p._p)
        c0._tc.append(mkpara(tpl_time,tm))
        c1._tc.append(mkpara(tpl_title,lines[0]))
        for l in lines[1:]: c1._tc.append(mkpara(tpl_body,l))
    sp=strip_bm(copy.deepcopy(tpl_head)); sect.addprevious(sp); setp(Paragraph(sp,None),'')
d.save('out_別紙１_プログラム_ACDR2026_案_0928.docx')
# ================= D. 別紙２ 過去の開催地 =================
d=Document('a_r_acc.docx'); P=[p for p in d.paragraphs if p.text.strip()]
ins_after(P[-1],'2025年12月　東京')
d.save('out_別紙２_過去のアジア防災会議開催地_0928.docx')
print('built')
