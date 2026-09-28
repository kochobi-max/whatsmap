from docx import Document
from docx.text.paragraph import Paragraph
import copy
def strip_bm(e):
    for x in list(e.iter()):
        if x.tag.endswith('}bookmarkStart') or x.tag.endswith('}bookmarkEnd'): x.getparent().remove(x)
    return e
def setp(p,text):
    for h in list(p._p):
        if h.tag.endswith('}hyperlink'): p._p.remove(h)
    runs=p.runs
    keep=[r for r in runs if r._r.xpath('.//w:drawing')]
    tr=[r for r in runs if r not in keep]
    if tr:
        tr[0].text=text
        for r in tr[1:]: r._r.getparent().remove(r._r)
    else: p.add_run(text)
def ins_after(p,text):
    e=strip_bm(copy.deepcopy(p._p)); p._p.addnext(e); q=Paragraph(e,p._parent); setp(q,text); return q
d=Document('b_2025_acc.docx'); P=d.paragraphs
setp(P[0],'令和８年10月　日')
setp(P[6],'「アジア防災会議2026」をシンガポール共和国で開催\n～アジア各国の防災能力の向上とアジア地域の防災ネットワークの充実・強化を図ります～')
setp(P[9],'内閣府政策統括官（防災担当）は、一般財団法人アジア防災センター（ADRC）と共同して、毎年「アジア防災会議」を開催しています。今回はシンガポール共和国において、シンガポール民間防衛庁（SCDF）を共同ホストとして、シンガポール国際防災・危機管理エキスポ（SIDEX）との合同開催「SIDEX-ACDR2026」の中で開催しますのでお知らせします。')
setp(P[10],'なお、ACDR2026ウェビナー（URLは確定後に記載）より、どなたでも会議を御覧いただけます。')
P[11]._p.getparent().remove(P[11]._p)
setp(P[13],'＜アジア防災会議2026の概要＞')
setp(P[15],'１　日程　　令和８年10月28日（水）、29日（木）　※プログラムは別紙ご参照')
ins_after(P[15],'　　　　　　　　（SIDEX-ACDR2026全体は10月27日（火）から30日（金）まで）')
setp(P[16],'２　場所  　シンガポールEXPO（シンガポール共和国）')
setp(P[18],'　　　　　　　　（ACDR2026ウェビナー（URLは確定後に記載））')
setp(P[19],'３　主催　　内閣府政策統括官（防災担当）、一般財団法人アジア防災センター（ADRC）、')
q=ins_after(P[19],'　　　　　　　　シンガポール民間防衛庁（SCDF）（共同ホスト）')
ins_after(q,'　　　　　　　　※SIDEX（主催：COSEM）との合同開催「SIDEX-ACDR2026」として実施')
setp(P[22],'５　使用言語　　英語　')
setp(P[23],'６　特設HP　　https://acdr.adrc.asia/home/acdr2026　※会議の詳細内容はこちら')
P[27]._p.getparent().remove(P[27]._p)
P[30]._p.getparent().remove(P[30]._p)
P[28]._p.getparent().remove(P[28]._p)
P[12]._p.getparent().remove(P[12]._p)
d.save('out3_内閣府記者発表資料_ACDR2026_案_0928.docx'); print('built3')
