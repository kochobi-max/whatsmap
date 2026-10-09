from docx import Document
from docx.text.paragraph import Paragraph
import copy
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
def strip_bm(e):
    for x in list(e.iter()):
        if x.tag.endswith('}bookmarkStart') or x.tag.endswith('}bookmarkEnd'): x.getparent().remove(x)
    return e
def setp(p,text):
    for h in list(p._p):
        if h.tag.endswith('}hyperlink'): p._p.remove(h)
    rs=p.runs
    if not rs: p.add_run(text); return
    rs[0].text=text
    for r in rs[1:]: r._r.getparent().remove(r._r)
def ins_after(p,text):
    e=strip_bm(copy.deepcopy(p._p)); p._p.addnext(e); q=Paragraph(e,p._parent); setp(q,text); return q
def fix_txbx(d, reps):
    for tb in d.element.body.iter(W+'txbxContent'):
        for p in tb.iter(W+'p'):
            ts=list(p.iter(W+'t'))
            if not ts: continue
            full=''.join(t.text or '' for t in ts); new=full
            for a,b in reps: new=new.replace(a,b)
            if new!=full:
                ts[0].text=new; ts[0].set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                for t in ts[1:]: t.text=''
THEME='先端技術で未来に備える越境的な気候・災害リスク削減 ― 強靭な社会に向けて'
# ================= ADRC 記者発表（資料配付）: 2024 様式 =================
d=Document('f_adrc2024_acc.docx'); P=d.paragraphs
r=d.tables[0].rows[2]
def setcell(c,text):
    ps=c.paragraphs; setp(ps[0],text)
    for p in ps[1:]: p._p.getparent().remove(p._p)
setcell(r.cells[0],'10/　(　)\n　：00'); setcell(r.cells[3],'理事長　小川　雄二郎')
setp(P[3],'「アジア防災会議2026」の開催について')
setp(P[6],'一般財団法人アジア防災センター（ADRC）は、アジア各国の防災能力の向上及びアジア地域での防災ネットワークの充実・強化を図るため、毎年、日本国政府（内閣府）等との共催により「アジア防災会議（ACDR）」を開催しています。今回はシンガポール共和国において、シンガポール民間防衛庁（SCDF）を共同ホストとして、シンガポール国際防災・危機管理エキスポ（SIDEX）との合同開催「SIDEX-ACDR2026」の中で、「'+THEME+'」をテーマに、下記のとおり「アジア防災会議2026」を開催します。')
setp(P[7],'なお、ACDR2026ウェビナー（URLは確定後に記載）より、どなたでも会議をご覧いただけます。')
setp(P[11],'１　開 催 日　　2026年10月28日（水）、29日（木）')
ins_after(P[11],'　　　　　　　　（SIDEX-ACDR2026全体は10月27日（火）から30日（金）まで）')
setp(P[12],'２　場　　所  　シンガポールEXPO（シンガポール共和国）')
setp(P[14],'（ACDR2026ウェビナー（URLは確定後に記載））')
setp(P[16],'３　主　　催　　内閣府（防災担当）、一般財団法人アジア防災センター（ADRC）、')
setp(P[17],'　　　　　　　　シンガポール民間防衛庁（SCDF）（共同ホスト）')
ins_after(P[17],'　　　　　　　　※SIDEX（主催：COSEM）との合同開催「SIDEX-ACDR2026」として実施')
setp(P[21],'　〇10月28日（水）　')
setp(P[22],'・合同開会式　　　シンガポール政府代表、日本政府代表　他')
setp(P[23],'・開会挨拶　　　　日本政府代表、シンガポール民間防衛庁長官、アジア防災センター長')
setp(P[24],'・ラウンドテーブル（第１部）　仙台防災枠組の実施状況と残る課題に関するメンバー国からの発表')
setp(P[25],'・アジア防災センター運営委員会（メンバー国のみ）')
setp(P[26],'　〇10月29日（木）　')
setp(P[27],'・ラウンドテーブル（第２部）　仙台防災枠組の実施状況と残る課題に関するメンバー国からの発表')
q=ins_after(P[27],'・テクニカルセッション　孤立地域・アクセス困難地域の防災DX：先端技術によるコミュニティ・レジリエンスの再定義')
q=ins_after(q,'・フィールドトリップ　シンガポール民間防衛アカデミー（CDA）')
q=ins_after(q,'　〇10月30日（金）　')
q=ins_after(q,'・SIDEX-ACDR2026会議（基調講演、パネルディスカッション）')
setp(P[28],'６　使用言語　　英語')
setp(P[29],'７　Ｈ　　Ｐ　　https://acdr.adrc.asia/home/acdr2026')
fix_txbx(d,[('ACDR2024','ACDR2026'),('荒木田、中村','荒木田'),('ma-arakida@adrc.asia/an-nakamura@adrc.asia','ma-arakida@adrc.asia'),('https://acdr.adrc.asia','https://acdr.adrc.asia/home/acdr2026')])
d.save('out2_記者発表（資料配付）_ACDR2026_案_0928.docx')
# ================= 内閣府 記者発表: 2024 様式 =================
d=Document('g_cao2024_acc.docx'); P=d.paragraphs
setp(P[0],'令和８年10月　日')
setp(P[6],'「アジア防災会議2026」をシンガポール共和国で開催\n～アジア各国の防災能力の向上とアジア地域の防災ネットワークの充実・強化を図ります～')
setp(P[9],'内閣府政策統括官（防災担当）は、一般財団法人アジア防災センター（ADRC）と共同して、毎年「アジア防災会議」を開催しています。今回はシンガポール共和国において、シンガポール民間防衛庁（SCDF）を共同ホストとして、シンガポール国際防災・危機管理エキスポ（SIDEX）との合同開催「SIDEX-ACDR2026」の中で開催しますのでお知らせします。なお、ACDR2026ウェビナー（URLは確定後に記載）より、どなたでも会議を御覧いただけます。')
setp(P[11],'＜アジア防災会議2026の概要＞')
setp(P[13],'１　日程　　令和８年10月28日（水）、29日（木）')
ins_after(P[13],'　　　　　　　　（SIDEX-ACDR2026全体は10月27日（火）から30日（金）まで）')
setp(P[14],'２　場所  　シンガポールEXPO（シンガポール共和国）')
setp(P[16],'　　　　　　　　（ACDR2026ウェビナー（URLは確定後に記載））')
setp(P[17],'３　主催　　内閣府政策統括官（防災担当）、一般財団法人アジア防災センター（ADRC）、')
setp(P[18],'　　　　　　　　シンガポール民間防衛庁（SCDF）（共同ホスト）')
ins_after(P[18],'　　　　　　　　※SIDEX（主催：COSEM）との合同開催「SIDEX-ACDR2026」として実施')
setp(P[22],'６　特設HP　 https://acdr.adrc.asia/home/acdr2026　※会議の詳細内容はこちら')
fix_txbx(d,[('荒木田、中村','荒木田'),('ma-arakida@adrc.asia, an-nakamura@adrc.asia,','ma-arakida@adrc.asia, acdr2026sgp@adrc.asia'),('担当：梅津','担当：〇〇')])
d.save('out2_内閣府記者発表資料_ACDR2026_案_0928.docx')
print('built2')
