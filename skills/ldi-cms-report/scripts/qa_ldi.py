#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
qa_ldi.py — 出来上がった docx を機械で検める（Step 8 のゲート）

    python scripts/qa_ldi.py <docxのパス> [--today YYYY-MM-DD]

終了コード
    0  PASS または WARN のみ（送ってよい。WARN はメール末尾に列挙する）
    1  FAIL が1件以上（**送らない**）
    2  実行できなかった

## なぜこれが要るか

SKILL.md Step 8 は `report-qa` の呼び出しを**送信前の BLOCKING ゲート**としている。
2026-10-05 に確認したところ、**呼び出し先のスクリプトが存在しなかった。**

    **無い  skills/report-qa/scripts/qa_report.py

つまり**毎朝のレポートは機械検証を1度も通っていない。**
ゲートがあるつもりで、無かった。`report-qa` スキルには SKILL.md も実体が無いため、
検証はこのスキルの中に持つ（外に置くと、また無いまま参照し続ける）。

## 何を見るか

このリポジトリでこれまでに実際に起きた事故の型に合わせてある。

  - 空欄のまま出力された（コロンビアで被害19行・リンク84行が全部空だった）
  - プレースホルダが残ったまま出力された
  - 件数の表記が中身と食い違った
  - 言語の取り違え（英語欄に日本語、日本語欄に英語）

**「ビルドが通った」は検証にならない。** いずれも例外を出さずに通った。
"""
import argparse
import re
import sys

try:
    from docx import Document
except ImportError:
    sys.stderr.write("STATUS: FAIL no-python-docx\npip install python-docx\n")
    sys.exit(2)

EXPECTED_LABELS = [
    "レポートの種類 / Report Type",
    "タイトル（日本語）",
    "タイトル（英語）/ Title",
    "内容（日本語）",
    "内容（英語）/ Outlines",
    "URL1",
    "URL2",
]
# 内容欄に入れてはいけないメタ情報（SKILL.md Step 7「内容欄に含めない」）
BANNED_IN_BODY = ["GLIDE", "GDACS", "Sentinel Asia", "センチネルアジア"]
PLACEHOLDERS = ["TBD", "TODO", "FIXME", "XXXX", "xxxx", "（未記入）", "placeholder"]
SOFT = ["確認中", "TBC"]

findings = []


def fail(msg):
    findings.append(("FAIL", msg))


def warn(msg):
    findings.append(("WARN", msg))


def has_japanese(s):
    return bool(re.search(r"[぀-ヿ一-鿿]", s))


def has_latin_words(s):
    return bool(re.search(r"[A-Za-z]{3,}", s))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    ap.add_argument("--today")
    a = ap.parse_args()

    try:
        doc = Document(a.docx)
    except Exception as e:
        print("STATUS: FAIL unreadable")
        print("  " + str(e))
        sys.exit(2)

    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    text_all = "\n".join(paras)

    # ---- 見出しとサブヘッダ ----
    if not any(p.startswith("ADRC LDI CMS レポート入力データ") for p in paras):
        fail("タイトル行「ADRC LDI CMS レポート入力データ」が無い")
    sub = next((p for p in paras if p.startswith("作成日:")), None)
    if not sub:
        fail("サブヘッダ（作成日:…）が無い")
    created = None
    if sub:
        m = re.search(r"作成日:\s*(\d{4}-\d{2}-\d{2})", sub)
        if not m:
            fail("サブヘッダに作成日（YYYY-MM-DD）が無い: " + sub[:60])
        else:
            created = m.group(1)
            if a.today and created != a.today:
                warn("作成日が今日ではない（%s、今日 %s）" % (created, a.today))

    # ---- エントリ表を集める（7行x2列でラベルが規定どおりのもの） ----
    entry_tables = []
    for t in doc.tables:
        if len(t.columns) == 2 and len(t.rows) == len(EXPECTED_LABELS):
            labels = [r.cells[0].text.strip() for r in t.rows]
            if labels == EXPECTED_LABELS:
                entry_tables.append(t)
                continue
            if labels[0].startswith("レポートの種類"):
                fail("エントリ表のラベルが規定と違う: " + " / ".join(labels))
        elif len(t.columns) == 2 and t.rows and t.rows[0].cells[0].text.strip().startswith("レポートの種類"):
            fail("エントリ表の行数が %d（規定は %d）" % (len(t.rows), len(EXPECTED_LABELS)))

    n = len(entry_tables)
    if n == 0:
        warn("エントリが0件。更新が無い日ならこれでよい（メール本文で「更新0件」と書くこと）")

    # ---- フッタの件数と中身の一致 ----
    foot = next((p for p in paras if p.startswith("--- 更新必要")), None)
    if not foot:
        fail("フッタ「--- 更新必要 X件 / スキップ Y件 ---」が無い")
    else:
        m = re.search(r"更新必要\s*(\d+)件\s*/\s*スキップ\s*(\d+)件", foot)
        if not m:
            fail("フッタの件数が読めない: " + foot)
        elif int(m.group(1)) != n:
            fail("フッタの更新必要 %s件 と、実際のエントリ %d件 が食い違う" % (m.group(1), n))
    if sub:
        m2 = re.search(r"更新必要\s*(\d+)件", sub)
        if m2 and int(m2.group(1)) != n:
            fail("サブヘッダの更新必要 %s件 と、実際のエントリ %d件 が食い違う" % (m2.group(1), n))

    # ---- 各エントリの中身 ----
    for i, t in enumerate(entry_tables, 1):
        v = {lab: t.rows[j].cells[1].text.strip() for j, lab in enumerate(EXPECTED_LABELS)}
        tag = "No.%d" % i

        if v["レポートの種類 / Report Type"] != "Report/Articles":
            fail("%s レポートの種類が「Report/Articles」でない: %r"
                 % (tag, v["レポートの種類 / Report Type"]))

        for lab in ["タイトル（日本語）", "タイトル（英語）/ Title",
                    "内容（日本語）", "内容（英語）/ Outlines", "URL1"]:
            if not v[lab] or v[lab] == "—":
                fail("%s %s が空（または —）" % (tag, lab))

        u1 = v["URL1"]
        if u1 and u1 != "—" and not re.match(r"^https?://", u1):
            fail("%s URL1 がURLでない: %r" % (tag, u1[:60]))
        u2 = v["URL2"]
        if u2 and u2 != "—" and not re.match(r"^https?://", u2):
            fail("%s URL2 がURLでも「—」でもない: %r" % (tag, u2[:60]))

        # 言語の取り違え
        if v["内容（日本語）"] and not has_japanese(v["内容（日本語）"]):
            fail("%s 内容（日本語）に日本語が1文字も無い" % tag)
        if v["内容（英語）/ Outlines"] and has_japanese(v["内容（英語）/ Outlines"]):
            fail("%s 内容（英語）に日本語が混ざっている" % tag)
        if v["内容（英語）/ Outlines"] and not has_latin_words(v["内容（英語）/ Outlines"]):
            fail("%s 内容（英語）に英単語が無い" % tag)

        # タイトルは「情報源名 + 日付」の形
        for lab in ["タイトル（日本語）", "タイトル（英語）/ Title"]:
            if v[lab] and not re.search(r"\d", v[lab]):
                warn("%s %s に日付らしい数字が無い（情報源名＋日付の形か確認）: %r"
                     % (tag, lab, v[lab][:40]))

        # 内容欄のメタ情報
        for lab in ["内容（日本語）", "内容（英語）/ Outlines"]:
            for b in BANNED_IN_BODY:
                if b in v[lab]:
                    fail("%s %s に「%s」が入っている（内容欄にメタ情報を含めない）" % (tag, lab, b))

        # 短すぎる本文
        for lab in ["内容（日本語）", "内容（英語）/ Outlines"]:
            if 0 < len(v[lab]) < 40:
                warn("%s %s が %d字。被害と対応が書けているか確認" % (tag, lab, len(v[lab])))

        # 発表機関名から書き出しているか（SKILL.md「発表機関名から書き出す」）
        if v["内容（日本語）"] and re.match(r"^[0-9０-９]", v["内容（日本語）"]):
            warn("%s 内容（日本語）が数字で始まっている（発表機関名から書き出す）" % tag)

    # ---- 文書全体の走査 ----
    cells_text = []
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                cells_text.append(c.text)
    whole = text_all + "\n" + "\n".join(cells_text)
    for ph in PLACEHOLDERS:
        if ph in whole:
            fail("プレースホルダ「%s」が残っている" % ph)
    for s in SOFT:
        if s in whole:
            warn("「%s」が残っている（意図的なら可。人が判断する）" % s)

    # ---- まとめ ----
    print("── LDI レポートのQA  %s" % a.docx)
    print("   エントリ %d件 / 作成日 %s" % (n, created or "?"))
    nf = sum(1 for lv, _ in findings if lv == "FAIL")
    nw = sum(1 for lv, _ in findings if lv == "WARN")
    for lv, msg in findings:
        print("   %s %s" % ("✗" if lv == "FAIL" else "!", msg))
    if not findings:
        print("   ✓ 指摘なし")
    print("SUMMARY: FAIL %d / WARN %d" % (nf, nw))
    if nf:
        print("STATUS: FAIL  **メールを送らない。** 上の ✗ を直してから作り直す。")
        sys.exit(1)
    if nw:
        print("STATUS: WARN  送ってよい。**WARN をメール本文の末尾に列挙する。**")
    else:
        print("STATUS: PASS")
    sys.exit(0)


if __name__ == "__main__":
    main()
